"""Regression tests for incomplete scraping and multi-city searches, offline
(a simulated city of businesses; no Serper credits, no network):
  1) a city search reaches a target of 100 when the city has enough
     businesses - phrases whose first page only repeats other phrases' firms
     are still read further (deep tier probe)
  2) Noida, then Prayagraj, Uttar Pradesh, India: each search sends only its
     own city, from the right Google country (gl=in), counts only its own
     records, keeps its own planner memory and Excel heading
  3) repeated searches: the same target again finishes at once without
     spending credits; a higher target keeps searching (searches the first
     run never sent are not lost)
  4) empty results: exhausted, with the real reason (no results at all)
  5) a site whose record was refused only because the target was full, or
     whose fetch failed transiently, can be crawled by a later search
  6) requested / discovered / valid / duplicate / rejected / exported counts
     are kept separately and add up
Run:  python tests/test_bugfix_scrape_counts.py  (from the app/ directory)"""

from __future__ import annotations

import hashlib
import os
import sys
from _scratch import scratch_dir  # noqa: E402  (tests/_scratch.py)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend import config  # noqa: E402

_tmp = scratch_dir(prefix="scrape_counts_test_")
config.OUTPUT_DIR = _tmp
config.STATE_PATH = os.path.join(_tmp, "state.json")
config.EMAIL_VERIFY = False

from backend.collector import engine  # noqa: E402
from backend.collector.discovery import DiscoveryPlanner, SearchRequest  # noqa: E402
from backend.collector.engine import CollectionJob, StateStore, _transient  # noqa: E402
from backend.collector.exporter import section_row_counts  # noqa: E402
from backend.collector.geo import country_code, state_of_city  # noqa: E402
from backend.collector.search import SearchProvider, SerperProvider  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def h(*parts) -> int:
    return int(hashlib.md5("|".join(map(str, parts)).encode()).hexdigest()[:8], 16)


class City:
    """`n` businesses per city. Every phrase ranks them by popularity with a
    little phrase-specific noise, like Google Maps: page 1 of each phrase is
    mostly the same popular firms, later pages differ."""

    def __init__(self, n=300, noise=150):
        self.n, self.noise = n, noise

    def page(self, city, phrase, page):
        order = sorted(range(self.n), key=lambda i: i + h(city, phrase, i) % self.noise)
        pin = {"Noida": 201301, "Prayagraj": 211001}.get(city, 100001)
        return [{"title": f"{city} Legal Associates {i:03d}", "cid": f"{city}-{i}",
                 "phoneNumber": f"+91 9{h(city, i) % 10 ** 9:09d}",
                 "address": f"Sector {i % 60}, {city}, Uttar Pradesh {pin}, India",
                 "category": "Law firm"}
                for i in order[(page - 1) * 10: page * 10]]


class CityProvider(SearchProvider):
    name = "fake"
    supports_places = True

    def __init__(self, world, empty=False):
        self.world, self.empty = world, empty
        self.credits_used = 0
        self.calls: list[tuple[str, str, int]] = []
        self.gl = ""
        self.gls: set[str] = set()

    def _city(self, query):
        q = query.lower()
        return next((c for c in ("Noida", "Prayagraj") if c.lower() in q), "Elsewhere")

    def places_page(self, query, page):
        self.credits_used += 1
        self.calls.append(("places", query, page))
        self.gls.add(self.gl)
        if self.empty:
            return []
        return self.world.page(self._city(query), query.lower().split(" in ")[0], page)

    def organic(self, query, max_results):
        return self.organic_page(query, 1)

    def organic_page(self, query, page):
        self.credits_used += 1
        self.calls.append(("organic", query, page))
        self.gls.add(self.gl)
        return []

    def balance(self):
        return 10_000


class NoFetch:
    def fetch_raw(self, url):
        return None, url, "connection/DNS error"


def run(state, city, target, world=None, empty=False, category="Law_Firms", keywords=()):
    geo = {"city": city, "state": "Uttar Pradesh", "country": "India"}
    job = CollectionJob(state, category, list(keywords), f"{city}, Uttar Pradesh, India",
                        target, "serper", 12, geo=geo)
    job.fetcher = NoFetch()
    prov = CityProvider(world or City(), empty=empty)
    real = engine.make_provider
    engine.make_provider = lambda *_: prov
    try:
        job._run()
    finally:
        engine.make_provider = real
    return job, prov


def check_counts_add_up(job):
    c = job.counters
    assert c["discovered"] == c["valid"] + c["duplicates"] + c["rejected"] + c["over_target"], c


def test_planner_probes_cells_without_new_results():
    pl = DiscoveryPlanner(["lawyers", "advocates"], ["Noida, Uttar Pradesh, India"],
                          places=True, max_page=3)
    a = pl.next_requests(want_organic=False)[0]
    b = pl.next_requests(want_organic=False)[0]
    pl.on_result(a, results=10, new=8)
    pl.on_result(b, results=10, new=0)      # page 1 = firms the first phrase found
    while pl.escalate():
        pass                                # deep tier
    seen = []
    while True:
        reqs = pl.next_requests(want_organic=True)
        if not reqs:
            break
        r = reqs[0]
        seen.append((r.query.split(" in ")[0], r.kind, r.page))
        # the productive cell keeps yielding; the other one's probe finds nothing
        pl.on_result(r, results=10, new=5 if r.query.startswith("lawyers") else 0)
    assert ("advocates", "places", 2) in seen, f"zero-new cell never read further: {seen}"
    assert ("advocates", "organic", 1) in seen, seen
    assert not any(q == "advocates" and k == "places" and p > 2 for q, k, p in seen), \
        f"a probe that found nothing must not page on: {seen}"
    assert ("lawyers", "places", 5) in seen, seen
    print("  deep tier: a cell whose page 1 only repeated others gets one probe page")


def test_unsent_search_is_handed_back():
    geos = ["Noida, Uttar Pradesh, India"]
    pl = DiscoveryPlanner(["lawyers", "advocates"], geos, places=True)
    first = pl.next_requests(want_organic=True)        # places + organic, cell 1
    second = pl.next_requests(want_organic=True)       # cell 2
    for r in first:
        pl.on_result(r, results=10, new=9)
    for r in second:                                   # target reached meanwhile
        pl.on_unsent(r)
    saved = pl.export()
    again = DiscoveryPlanner(["lawyers", "advocates"], geos, saved=saved, places=True)
    nxt = again.next_requests(want_organic=True)
    assert nxt and nxt[0].query == second[0].query, \
        f"never-sent search lost for later runs: {[r.query for r in nxt]}"
    assert first[0].query not in {r.query for r in nxt}
    # a never-sent page 2 does not mark the cell's result list as ended
    fu = SearchRequest("places", first[0].query, 2, first[0].phrase, first[0].geo,
                       meta={"parent_rate": 0.9})
    pl.used.add(f"{pl._cell_key(fu.phrase, fu.geo)}#p2")
    pl.on_unsent(fu)
    cell = pl.cells[pl._cell_key(fu.phrase, fu.geo)]
    assert cell["pp"] == 1 and cell["pn"] == 10, cell
    print("  searches never sent go back to the planner; page memory untouched")


def test_noida_then_prayagraj_100_each():
    state = StateStore()
    noida, pn = run(state, "Noida", 100)
    assert noida.status == "completed", (noida.status, noida.stop_reason)
    assert noida.category_count() == 100, noida.category_count()
    assert {r["Search Location"] for r in state.records["Law_Firms"]} == \
        {"Noida, Uttar Pradesh, India"}
    check_counts_add_up(noida)

    prayag, pp = run(state, "Prayagraj", 100)
    assert prayag.status == "completed", (prayag.status, prayag.stop_reason)
    assert prayag.category_count() == 100 and prayag.snapshot(False)["collected"] == 100
    assert prayag.counters["valid"] == 100, "Prayagraj counted Noida's records"
    check_counts_add_up(prayag)
    # its own filters: only its own city in every query, from India
    assert all("Prayagraj, Uttar Pradesh, India" in q for _, q, _ in pp.calls), pp.calls[:5]
    assert not any("noida" in q.lower() for _, q, _ in pp.calls)
    assert pn.gls == {"in"} and pp.gls == {"in"}, (pn.gls, pp.gls)
    recs = state.records["Law_Firms"]
    by_loc = {}
    for r in recs:
        by_loc[r["Search Location"]] = by_loc.get(r["Search Location"], 0) + 1
        assert r["City"] in ("Noida", "Prayagraj")
        assert r["City"] in r["Search Location"], r
    assert by_loc == {"Noida, Uttar Pradesh, India": 100,
                      "Prayagraj, Uttar Pradesh, India": 100}, by_loc
    assert set(k for k in state.discovery if k.startswith("Law_Firms@@")) == {
        "Law_Firms@@noida, uttar pradesh, india",
        "Law_Firms@@prayagraj, uttar pradesh, india"}
    # Excel: the state's sheet, one heading per city, every record exported
    path = os.path.join(config.OUTPUT_DIR, state.all_categories()["Law_Firms"]["file"])
    wb = load_workbook(path, read_only=True)
    assert wb.sheetnames == ["Uttar Pradesh"], wb.sheetnames
    header = next(wb["Uttar Pradesh"].iter_rows(max_row=1, values_only=True))
    assert header[0] == "Company Name" and "Record Status" in header
    wb.close()
    assert section_row_counts(path) == {"Noida, Uttar Pradesh, India": 100,
                                        "Prayagraj, Uttar Pradesh, India": 100}
    assert noida.counters["exported"] == 100 and prayag.counters["exported"] == 100
    last = list(prayag.log_lines)[-1]
    assert "requested 100, collected 100" in last and "100 rows exported under " \
        "'Prayagraj, Uttar Pradesh, India'" in last, last
    print("  Noida then Prayagraj: 100 + 100, own queries / gl=in / counts / heading")
    return state


def test_repeated_searches(state):
    again, prov = run(state, "Noida", 100)
    assert again.status == "completed" and prov.credits_used == 0, \
        (again.status, prov.credits_used)
    assert "already has 100 records" in "\n".join(again.log_lines)
    more, prov = run(state, "Noida", 120)
    assert prov.credits_used > 0, "a higher target must search again"
    assert more.category_count() > 100, more.category_count()
    assert more.status in ("completed", "exhausted"), more.status
    if more.status == "exhausted":
        assert "searches returned" in more.stop_reason, more.stop_reason
    assert more.category_count() == 100 + more.counters["valid"]
    # Prayagraj is untouched by Noida's repeated searches
    assert state.location_count("Law_Firms", "prayagraj, uttar pradesh, india") == 100
    print(f"  repeated searches: same target -> 0 credits; target 120 -> "
          f"{more.category_count()} ({more.status})")


def test_empty_results():
    state = StateStore()
    job, prov = run(state, "Noida", 100, empty=True)
    assert job.status == "exhausted" and job.category_count() == 0
    assert "returned no results at all" in job.stop_reason, job.stop_reason
    assert "already found" not in job.stop_reason and "already collected" not in job.stop_reason
    assert job.counters["exported"] == 0
    print("  empty results: exhausted at 0 with the real reason")


def test_small_city_reports_shortfall_honestly():
    state = StateStore()
    job, _ = run(state, "Noida", 100, world=City(n=40, noise=10))
    assert job.status == "exhausted" and job.category_count() == 40, \
        (job.status, job.category_count())
    assert "40 stored" in job.stop_reason, job.stop_reason
    last = list(job.log_lines)[-1]
    assert "Could only find 40 of 100" in last and "requested 100, collected 40" in last
    print("  a city with 40 businesses: all 40 collected, shortfall explained")


def test_crawl_memory_released():
    state = StateStore()
    loc = {"city": "Noida", "state": "Uttar Pradesh", "country": "India"}
    job = CollectionJob(state, "Law_Firms", [], "Noida, Uttar Pradesh, India", 1,
                        "serper", 4, geo=loc)

    def rec(name, site):
        return job._base_record("q") | {"Company Name": name, "Official Website": site,
                                        "Business Phone": f"+91 98{h(name) % 10 ** 8:08d}",
                                        "_ck": site.split("//")[1]}
    for key in ("alpha.example", "beta.example"):
        assert state.mark_crawled(key, "Law_Firms")
    assert job._try_add(rec("Alpha Law Chambers", "https://alpha.example"))
    assert not job._try_add(rec("Beta Law Chambers", "https://beta.example"))
    assert job.counters["over_target"] == 1
    assert "alpha.example" in state.crawled, "a stored record's site stays crawled"
    assert "beta.example" not in state.crawled, "refused for a full target: crawl again later"
    assert _transient("timeout") and _transient("HTTP 503") and _transient("HTTP 429")
    assert _transient("connection/DNS error")
    assert not _transient("HTTP 404") and not _transient("invalid SSL certificate")
    assert not _transient("blocked (private address)")
    print("  crawl memory: refused-for-target and transient failures are released")


def test_geo_and_gl():
    assert country_code("India") == "in" and country_code("United States") == "us"
    assert country_code("UK") == "gb" and country_code("Atlantis") == ""
    assert state_of_city("Prayagraj", "India") == "Uttar Pradesh"
    # no state typed: Prayagraj still goes to the Uttar Pradesh sheet
    job = CollectionJob(StateStore(), "Law_Firms", [], "Prayagraj, India", 5, "serper", 4,
                        geo={"city": "Prayagraj", "state": "", "country": "India"})
    assert job.search_place["Search State"] == "Uttar Pradesh", job.search_place
    assert job.search_place["Search Location"] == "Prayagraj, Uttar Pradesh, India"

    class Session:
        def __init__(self):
            self.payloads = []

        def post(self, url, json=None, timeout=None, headers=None):
            self.payloads.append(json)

            class R:
                status_code = 200
                headers = {}
                text = "{}"

                @staticmethod
                def json():
                    return {"places": [], "organic": []}
            return R()
    p = SerperProvider("test-key")
    p.session = Session()
    p.places_page("law firms in Noida", 1)
    p.gl = "in"
    p.places_page("law firms in Noida", 2)
    p.organic("law firms in Noida", 10)
    assert "gl" not in p.session.payloads[0], "no country set: Serper's default"
    assert all(pl.get("gl") == "in" for pl in p.session.payloads[1:]), p.session.payloads
    print("  gl from the searched country; Prayagraj maps to Uttar Pradesh")


def main():
    test_planner_probes_cells_without_new_results()
    test_unsent_search_is_handed_back()
    state = test_noida_then_prayagraj_100_each()
    test_repeated_searches(state)
    test_empty_results()
    test_small_city_reports_shortfall_honestly()
    test_crawl_memory_released()
    test_geo_and_gl()
    print("\nALL SCRAPE COUNT TESTS PASSED")


if __name__ == "__main__":
    main()
