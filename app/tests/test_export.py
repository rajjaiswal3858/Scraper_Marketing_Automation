"""Excel export: the file opens cleanly and its cells are usable - phone
numbers keep their "+", the score is a number, the time a real date, links
are clickable, empty values are empty cells, and the formula guard still
neutralizes injected formulas.
Run:  python tests/test_export.py  (from the app/ directory)"""

from __future__ import annotations

import os
import sys
from _scratch import scratch_dir  # noqa: E402  (tests/_scratch.py)
import zipfile
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from backend.collector.exporter import COLUMNS, write_category_file  # noqa: E402


def main():
    out = scratch_dir(prefix="export_test_")
    records = [
        {"Company Name": "Acme Tax LLP", "Category": "Finance",
         "Official Website": "https://acme.example", "Business Email": "info@acme.example",
         "Business Phone": "+91 98765 43210", "Country": "India",
         "Source URL": "https://acme.example/contact", "Collected At": "2026-10-12T21:58:48",
         "Confidence Score": "85"},
        {"Company Name": "=HYPERLINK(\"http://evil\",\"x\")", "Category": "Finance",
         "Business Phone": "+SUM(1,1)", "Services": "@cmd", "Business Email": "not an email",
         "Official Website": "javascript:alert(1)", "Confidence Score": "",
         "Collected At": "yesterday"},
    ]
    write_category_file("Finance", "01_Finance.xlsx", records, out)
    path = os.path.join(out, "01_Finance.xlsx")
    assert zipfile.ZipFile(path).testzip() is None
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == COLUMNS
    col = {name: i for i, name in enumerate(COLUMNS)}
    # row 2 is the location heading (these records predate location tags)
    assert ws["A2"].value.startswith("Search location: Earlier results")
    good, evil = list(ws[3]), list(ws[4])

    assert good[col["Business Phone"]].value == "+91 98765 43210"
    print("  phone numbers keep their leading +")
    assert good[col["Confidence Score"]].value == 85
    print("  confidence score is a number")
    when = good[col["Collected At"]]
    assert when.value == datetime(2026, 10, 12, 21, 58, 48) and "yy" in when.number_format
    print("  collected-at is a real date")
    assert good[col["Official Website"]].hyperlink.target == "https://acme.example"
    assert good[col["Business Email"]].hyperlink.target == "mailto:info@acme.example"
    print("  website / email are clickable")
    assert good[col["Subcategory"]].value is None
    print("  missing values are empty cells")
    assert ws.auto_filter.ref == "A1:W4" and ws.freeze_panes == "A2"
    print("  header has filter arrows and stays frozen")

    for name in ("Company Name", "Business Phone", "Services"):
        v = evil[col[name]].value
        assert isinstance(v, str) and not v.startswith(("=", "+", "@")), (name, v)
        assert evil[col[name]].data_type != "f"
    assert evil[col["Official Website"]].hyperlink is None
    assert evil[col["Business Email"]].hyperlink is None
    assert evil[col["Collected At"]].value == "yesterday"
    print("  injected formulas neutralized; no links to non-web targets")

    location_layout(out)
    print("ALL EXPORT TESTS PASSED")


def rec(name, query="", loc="", state="", country="USA"):
    r = {"Company Name": name, "Category": "Law Firms", "Business Email": f"info@{name.lower()}.com",
         "Official Website": f"https://{name.lower()}.com", "Search Query": query,
         "Collected At": "2026-10-16T10:00:00", "Confidence Score": "80"}
    if loc:
        r.update({"Search Location": loc, "Search State": state, "Search Country": country})
    return r


def location_layout(out):
    """State sheet, a heading per searched location in search order, later
    searches appended below, every earlier record kept unchanged."""
    from backend.collector.exporter import read_data_rows
    old = [rec("OldHouston", "law firms in Houston, Texas"),      # before location tags:
           rec("OldTexas", "law firms in Texas, USA")]              # an older Texas search
    ca = [rec(f"CaFirm{i}", "law firms in San Diego, California", "California, USA",
              "California") for i in range(3)]
    path = os.path.join(out, "05_Law_Firms.xlsx")
    write_category_file("Law Firms", "05_Law_Firms.xlsx", old + ca, out)
    before = read_data_rows(path)

    la = [rec(f"LaFirm{i}", "law firms in Los Angeles, California",
              "Los Angeles, California, USA", "California") for i in range(2)]
    tx = [rec("TxFirm", "law firms in Austin, Texas", "Texas, USA", "Texas")]
    write_category_file("Law Firms", "05_Law_Firms.xlsx", old + ca + la + tx, out)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Texas", "California"], wb.sheetnames

    cal = [r[0] for r in wb["California"].iter_rows(values_only=True)]
    assert cal[0] == "Company Name"
    assert cal[1] == "Search location: California, USA   (3 records)", cal[1]
    assert cal[2:5] == ["CaFirm0", "CaFirm1", "CaFirm2"]
    assert cal[5] is None                                  # spacer
    assert cal[6] == "Search location: Los Angeles, California, USA   (2 records)", cal[6]
    assert cal[7:9] == ["LaFirm0", "LaFirm1"]
    print("  California sheet: state results first, Los Angeles appended below its own heading")

    tex = [r[0] for r in wb["Texas"].iter_rows(values_only=True)]
    assert tex[1].startswith("Search location: Earlier results") and tex[2:4] == ["OldHouston", "OldTexas"]
    assert tex[5].startswith("Search location: Texas, USA") and tex[6] == "TxFirm"
    print("  older state search: kept in its state's sheet, above the newer Texas search")

    # an older COUNTRY-wide search (many states, bare city names) stays together
    usa = [rec("UsaNY", "law firms in New York"), rec("UsaHou", "law firms in Houston, Texas"),
           rec("UsaLA", "law firms in Los Angeles")]
    write_category_file("Law Firms", "05b.xlsx", usa, out)
    wb2 = load_workbook(os.path.join(out, "05b.xlsx"))
    assert wb2.sheetnames == ["Law Firms"], wb2.sheetnames
    assert [r[0] for r in wb2.active.iter_rows(values_only=True)][2:5] == ["UsaNY", "UsaHou", "UsaLA"]
    print("  older country-wide search: kept together, in its original order")

    after = read_data_rows(path)
    assert len(after) == len(old + ca + la + tx)
    by_name = {r[0]: r for r in after}
    for r in before:
        assert by_name[r[0]] == r, f"earlier record changed: {r[0]}"
    print("  every earlier record is still there, field for field")


if __name__ == "__main__":
    main()
