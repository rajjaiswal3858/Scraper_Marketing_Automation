"""Custom search query tests WITHOUT Serper credits or network access.

A user-typed Google query is searched exactly as typed, first, and once; alone
it is the whole search, beside keywords the usual plan follows it; operator
queries skip Places; results go through the normal pipeline and export.
Run:  python tests/test_custom_query.py   (from the app/ directory)
"""

from __future__ import annotations

import os
import sys
from _scratch import scratch_dir  # noqa: E402  (tests/_scratch.py)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend import config  # noqa: E402

_tmp = scratch_dir(prefix="custom_query_test_")
config.OUTPUT_DIR = _tmp
config.STATE_PATH = os.path.join(_tmp, "demo_state.json")

from backend.collector import emailcheck, engine  # noqa: E402
from backend.collector.categories import (parse_custom_queries, query_terms,  # noqa: E402
                                          search_phrases)
from backend.collector.discovery import DiscoveryPlanner, has_operators  # noqa: E402
from backend.collector.engine import CollectionJob, StateStore  # noqa: E402
from backend.collector.exporter import read_data_rows  # noqa: E402
from backend.collector.search import SearchProvider  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

emailcheck.verify = lambda e: emailcheck.Verdict(e, "Valid", 100, "test: mailbox exists")

PAGES = {
    "http://ledgerpro.example": """
      <html><head><title>LedgerPro Accounting | CPA Firm</title></head>
      <body><a href="mailto:hello@ledgerpro.example">Email</a>
      <a href="tel:+15125550111">Call</a>
      <address>100 Congress Ave, Austin, Texas 78701, USA</address></body></html>""",
    "http://lonestartax.example": """
      <html><head><title>Lone Star Tax Advisors | Accounting Services</title></head>
      <body><a href="mailto:info@lonestartax.example">mail</a>
      <a href="tel:+17135550122">call</a>
      <address>500 Main St, Houston, Texas 77002, USA</address></body></html>""",
}


class FakeFetcher:
    def fetch_raw(self, url: str):
        html = PAGES.get(url.rstrip("/"))
        if html is None:
            return None, url, "connection/DNS error"
        return html.encode("utf-8"), url, ""


class RecordingProvider(SearchProvider):
    """Answers only the first search of each kind; records every query."""
    name = "fake"
    supports_places = True

    def __init__(self, organic_urls, places_rows):
        self._organic, self._places = organic_urls, places_rows
        self.calls: list[tuple[str, str]] = []
        self.credits_used = 0

    def organic(self, query, max_results):
        self.credits_used += 1
        self.calls.append(("organic", query))
        first = sum(k == "organic" for k, _ in self.calls) == 1
        return list(self._organic) if first else []

    def places(self, query, pages):
        self.credits_used += 1
        self.calls.append(("places", query))
        first = sum(k == "places" for k, _ in self.calls) == 1
        return list(self._places) if first else []


def run_job(state, keywords, custom, geo=None, target=5, organic=(), places=()):
    geo = geo or {}
    location = ", ".join(p for p in (geo.get("city"), geo.get("state"),
                                     geo.get("country")) if p)
    job = CollectionJob(state, "Finance", keywords, location, target, "serper", 4,
                        geo=geo, custom_queries=custom)
    job.fetcher = FakeFetcher()
    prov = RecordingProvider(list(organic), list(places))
    real_make = engine.make_provider
    engine.make_provider = lambda *_: prov
    try:
        job._run()
    finally:
        engine.make_provider = real_make
    return job, prov


def parsing_tests():
    qs = parse_custom_queries("accounting firms in Texas\n\n  dentists in Austin,  TX \n"
                              "ACCOUNTING FIRMS IN TEXAS\nsite:linkedin.com \"CFO\"", 5, 300)
    assert qs == ["accounting firms in Texas", "dentists in Austin, TX",
                  'site:linkedin.com "CFO"'], qs
    assert len(parse_custom_queries("\n".join(f"q{i}" for i in range(9)), 5, 300)) == 5
    assert parse_custom_queries("x" * 400, 5, 300) == ["x" * 300]
    assert parse_custom_queries("", 5, 300) == []
    terms = query_terms(['site:linkedin.com "accounting" OR bookkeeping'])
    assert "accounting" in terms and "bookkeeping" in terms
    assert not {"site", "linkedin", "com"} & set(terms), terms
    assert has_operators('site:linkedin.com cpa') and has_operators('"cpa firm" texas')
    assert has_operators("cpa -jobs") and has_operators("cpa OR accountant")
    assert not has_operators("accounting firms in Texas")
    assert not has_operators("e-commerce agencies in London")
    print("parsing tests passed")


def planner_tests():
    phrases = search_phrases("Finance", "Finance", ["accounting firms"])
    geos = ["Texas, USA", "Houston, Texas"]
    # custom query first, exactly as typed, before any phrase x location cell
    pl = DiscoveryPlanner(phrases, geos, places=True,
                          fixed_queries=["Email marketing agencies in London"])
    reqs = pl.next_requests(want_organic=True)
    assert {r.query for r in reqs} == {"Email marketing agencies in London"}
    assert {r.kind for r in reqs} == {"places", "organic"}
    nxt = pl.next_requests(want_organic=True)
    assert nxt[0].query == "accounting firms in Texas, USA", nxt[0].query
    # a custom query word-for-word equal to a generated cell is ONE search
    pl = DiscoveryPlanner(phrases, geos, places=True,
                          fixed_queries=["accounting  firms in texas, usa"])
    seen = []
    while True:
        reqs = pl.next_requests(want_organic=False)
        if not reqs:
            break
        if reqs[0].page == 1:          # further pages of a cell are not repeats
            seen.append(reqs[0].query.lower())
        pl.on_result(reqs[0], results=10, new=5)
    assert seen[0] == "accounting firms in texas, usa"
    assert sum(" ".join(q.split()) == "accounting firms in texas, usa" for q in seen) == 1, seen
    # operators: organic only (Places ignores them), also in the deep tier
    pl = DiscoveryPlanner([], [], places=True, fixed_queries=['site:clutch.co "cpa"'])
    reqs = pl.next_requests(want_organic=False, want_places=True)
    assert [r.kind for r in reqs] == ["organic"], reqs
    pl.on_result(reqs[0], results=10, new=8)
    while pl.escalate():
        pass
    deep = pl.next_requests(want_organic=True)
    assert deep and all(r.kind == "organic" for r in deep), deep
    assert deep[0].page == 2 and deep[0].query == 'site:clutch.co "cpa"'
    # persistence: a used custom query is never paid for again
    pl2 = DiscoveryPlanner([], [], saved=pl.export(), places=True,
                           fixed_queries=['site:clutch.co "cpa"'])
    assert pl2.candidate_cells() == 0 and pl2.next_requests(want_organic=True) == []
    print("planner tests passed")


def job_tests():
    places = [{"title": "LedgerPro Accounting", "website": "http://ledgerpro.example",
               "address": "100 Congress Ave, Austin, Texas 78701, USA",
               "phoneNumber": "+1 512-555-0111", "cid": "1"}]
    organic = ["http://lonestartax.example"]

    # custom query alone, no location: ONLY that query is searched, and the
    # results get their own heading
    state = StateStore()
    job, prov = run_job(state, [], ["accounting firms in Texas"],
                        organic=organic, places=places)
    assert {q for _, q in prov.calls} == {"accounting firms in Texas"}, prov.calls
    assert job.custom_only and job.planner.summary()["custom_queries"] == 1
    recs = state.records["Finance"]
    assert len(recs) == 2, [r.get("Company Name") for r in recs]
    for r in recs:
        assert r["Search Query"] == "accounting firms in Texas"
        assert r["Search Location"] == "Custom search: accounting firms in Texas"
        assert r["Business Email"] and r["Email Status"] == "Valid"
        assert r["Confidence Score"] >= 70, r
    snap = job.snapshot(include_log=False)
    assert snap["custom_queries"] == ["accounting firms in Texas"]
    assert "Add another custom query" in job.stop_reason, job.stop_reason

    # export: same workbook, same columns, the query is the section heading
    state.export_categories(["Finance"])
    xlsx = os.path.join(state.out_dir, "01_Finance.xlsx")
    wb = load_workbook(xlsx)
    cells = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value]
    assert any("Custom search: accounting firms in Texas" in str(v) for v in cells)
    rows = read_data_rows(xlsx)
    assert len(rows) == 2, rows

    # the same custom query again: already paid for, nothing re-searched,
    # nothing duplicated
    job2, prov2 = run_job(state, [], ["accounting firms in Texas"],
                          organic=organic, places=places, target=10)
    assert prov2.calls == [], prov2.calls
    assert len(state.records["Finance"]) == 2

    # custom query + keywords + location: custom query first, then the usual
    # keyword x location plan; a generated search identical to the custom
    # query is never sent twice
    state = StateStore()
    job, prov = run_job(state, ["accounting firms"], ["accounting firms in Texas, USA",
                                                      "cpa firms near Dallas"],
                        geo={"state": "Texas", "country": "USA"}, target=50,
                        organic=organic, places=places)
    queries = [q for _, q in prov.calls]
    assert queries[0] == "accounting firms in Texas, USA", queries[:4]
    assert "cpa firms near Dallas" in queries[:4], queries[:4]
    assert queries.count("accounting firms in Texas, USA") == 2   # places + organic, once
    assert any(q.endswith(", Texas") for q in queries), "keyword plan did not run"
    recs = state.records["Finance"]
    assert len(recs) == 2 and all(r["Search Location"] == "Texas, USA" for r in recs)

    # no custom query: behaviour unchanged (no custom cells, location heading)
    state = StateStore()
    job, prov = run_job(state, ["accounting firms"], [],
                        geo={"state": "Texas", "country": "USA"},
                        organic=organic, places=places)
    assert job.planner.fixed == [] and not job.custom_only
    assert prov.calls[0][1] == "accounting firms in Texas, USA"
    print("job tests passed")


# The marketing team's advanced queries, exactly as they type them.
ADVANCED = [
    'site:.com CPA in "726" "+1" "info" -za -uk -au -india -gov -edu',
    '"Health care" "Yuma AZ" "email" gmail',
    '"CPA in Texas" "email" "info" site:linkedin.com',
    'payroll AND "WA" inurl:contact email',
    'site:.com/contact consulting "Chicago" email -intitle:accounting',
]


def advanced_parsing_tests():
    from backend.collector.discovery import norm_query
    # every query survives the input box untouched, operators and all
    assert parse_custom_queries("\n".join(ADVANCED), 5, 300) == ADVANCED
    assert all(has_operators(q) for q in ADVANCED)
    # lower-case and / or are words, not operators: Places still runs
    assert not has_operators("law and tax firms in Austin")
    assert not has_operators("dentists or orthodontists in Austin")
    assert has_operators("cpa AROUND(3) texas") and has_operators("cpa | accountant")
    # Google's OR is case-sensitive: a different query, a different key
    assert norm_query("cpa OR tax") != norm_query("cpa or tax")
    assert norm_query("  CPA   firms ") == norm_query("cpa firms")
    # exclusions are never relevance terms
    terms = query_terms([ADVANCED[0], ADVANCED[4], 'cpa -"real estate"'])
    for bad in ("india", "gov", "edu", "accounting", "real estate", "intitle"):
        assert bad not in terms, (bad, terms)
    assert "cpa" in terms and "consulting" in terms and "chicago" in terms, terms
    print("advanced parsing tests passed")


class _Resp:
    def __init__(self, status, data):
        self.status_code, self._data, self.headers = status, data, {}
        self.text = "" if status == 200 else "temporarily unavailable"

    def json(self):
        return self._data


class WireSession:
    """Stands in for requests.Session: records the exact HTTP body requests
    would send and answers like Serper (one 503 first, to exercise retries)."""

    def __init__(self, fail_first=True):
        import requests
        self._requests = requests
        self.bodies: list[bytes] = []
        self.fail_first = fail_first

    def post(self, url, json=None, timeout=None, headers=None):
        req = self._requests.Request("POST", url, json=json, headers=headers).prepare()
        self.bodies.append(req.body)
        if self.fail_first:
            self.fail_first = False
            return _Resp(503, {})
        return _Resp(200, {"organic": [
            {"link": "https://www.linkedin.com/company/ledgerpro", "title": "LedgerPro",
             "snippet": "CPA in Texas. Email: info@ledgerpro.example"}],
            "credits": 1})


def provider_wire_tests():
    import json
    from backend.collector.search import SerperProvider
    p = SerperProvider("test-key")
    p.session = WireSession()
    for q in ADVANCED:
        p.session.fail_first = True
        urls = p.organic(q, 10)
        assert urls == ["https://www.linkedin.com/company/ledgerpro"]
        body = json.loads(p.session.bodies[-1].decode("utf-8"))
        assert body["q"] == q, (body, q)            # byte-exact after a retry
        assert p.session.bodies[-1] == p.session.bodies[-2]   # retry resent as-is
        assert "info@ledgerpro.example" in p.last_items[0]["snippet"]
    p.organic_page(ADVANCED[2], 4)
    body = json.loads(p.session.bodies[-1])
    assert body == {"q": ADVANCED[2], "num": 10, "page": 4}, body
    p.gl = "us"
    p.organic(ADVANCED[1], 10)
    assert json.loads(p.session.bodies[-1])["gl"] == "us"
    print("provider wire tests passed")


def deep_pagination_tests():
    # a custom query is read up to 10 pages while pages are full and new;
    # generated searches keep the old 3-page limit
    pl = DiscoveryPlanner(["accounting firm"], ["Austin, Texas"], places=False,
                          fixed_queries=[ADVANCED[3]])
    pages: dict[str, list[int]] = {}
    while True:
        reqs = pl.next_requests(want_organic=True, want_places=False)
        if not reqs:
            if not pl.escalate():
                break
            continue
        for r in reqs:
            pages.setdefault(r.query, []).append(r.page)
            pl.on_result(r, results=10, new=6)
    assert pages[ADVANCED[3]] == list(range(1, 11)), pages
    assert max(pages["accounting firm in Austin, Texas"]) == 3, pages
    print("deep pagination tests passed")


SNIPPET_PAGES = {
    # reached only through the email Google showed on its LinkedIn page
    "https://ledgerpro.example": """
      <html><head><title>LedgerPro Accounting | CPA Firm</title></head>
      <body><a href="mailto:hello@ledgerpro.example">Email</a>
      <a href="tel:+15125550111">Call</a></body></html>""",
    # its own site shows a phone but no email; Google's snippet shows one
    "https://yumacare.example/contact": """
      <html><head><title>Yuma Family Care | Health care clinic</title></head>
      <body><a href="tel:+19285550133">Call us</a>
      <address>1 Main St, Yuma, AZ 85364, USA</address></body></html>""",
}


class SnippetFetcher:
    def fetch_raw(self, url: str):
        html = SNIPPET_PAGES.get(url.rstrip("/"))
        if html is None:
            return None, url, "connection/DNS error"
        return html.encode("utf-8"), url, ""


class SnippetProvider(SearchProvider):
    """Organic only, with full items (title / snippet) like Serper's."""
    name = "fake"
    supports_places = True

    def __init__(self, items):
        self.items, self.calls, self.credits_used = items, [], 0
        self.last_items: list[dict] = []

    def organic(self, query, max_results):
        return self.organic_page(query, 1)

    def organic_page(self, query, page):
        self.credits_used += 1
        self.calls.append(("organic", query, page))
        self.last_items = list(self.items) if page == 1 else []
        return [i["link"] for i in self.last_items]

    def places(self, query, pages):
        self.calls.append(("places", query, 1))
        return []


def snippet_job_tests():
    items = [
        {"link": "https://www.linkedin.com/company/ledgerpro", "title": "LedgerPro",
         "snippet": "CPA in Texas ... email info@ledgerpro.example for an appointment"},
        {"link": "https://www.linkedin.com/in/someone", "title": "Jane, CPA",
         "snippet": "Reach me at jane.cpa.texas@gmail.com"},          # no website
        {"link": "https://www.facebook.com/x", "title": "x", "snippet": "no email"},
        {"link": "https://yumacare.example/contact", "title": "Yuma Family Care",
         "snippet": "Email: frontdesk@yumacare.example, other@unrelated.example"},
    ]
    for q in (ADVANCED[2], ADVANCED[1]):
        state = StateStore()
        job = CollectionJob(state, "Finance", [], "", 5, "serper", 4, custom_queries=[q])
        job.fetcher = SnippetFetcher()
        prov = SnippetProvider(items)
        real_make = engine.make_provider
        engine.make_provider = lambda *_: prov
        try:
            job._run()
        finally:
            engine.make_provider = real_make
        # Places is never asked (operators); the query went out exactly as typed
        assert {c[0] for c in prov.calls} == {"organic"}, prov.calls
        assert all(c[1] == q for c in prov.calls), prov.calls
        recs = {r["Official Website"]: r for r in state.records["Finance"]}
        assert set(recs) == {"https://ledgerpro.example", "https://yumacare.example"}, recs
        lp = recs["https://ledgerpro.example"]
        assert lp["Business Email"] == "hello@ledgerpro.example"   # the site's own
        yc = recs["https://yumacare.example"]
        # the site's own address from the snippet; never the unrelated one
        assert yc["Business Email"] == "frontdesk@yumacare.example", yc
        assert yc["Email Source"] == "https://yumacare.example/contact"
        assert all(r["Search Query"] == q for r in recs.values())
        assert job.counters["snippet_sites"] == 1 and job.counters["snippet_emails"] == 1
        # through the usual export: same workbook layout
        state.export_categories(["Finance"])
        rows = read_data_rows(os.path.join(state.out_dir, "01_Finance.xlsx"))
        assert len(rows) == 2, rows
    print("snippet job tests passed")


def main():
    parsing_tests()
    advanced_parsing_tests()
    provider_wire_tests()
    planner_tests()
    deep_pagination_tests()
    job_tests()
    snippet_job_tests()
    print("ALL CUSTOM QUERY TESTS PASSED")


if __name__ == "__main__":
    main()
